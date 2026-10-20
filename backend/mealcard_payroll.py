"""
mealcard_payroll.py - turn the Paytronix "Stored Value Sales and Added Value
Detail Report" into the two workbooks payroll needs.

Ported from the desktop script make_payroll_report.py. The differences are all
about living in a web app instead of a folder:

  * input arrives as bytes (an upload), not a path to glob for in ~/Downloads
  * output is returned as bytes, not saved next to the script
  * the card -> name and card -> employee-ID lookups come from mealcard_db,
    not from a spreadsheet and a CSV sitting beside the code

The report logic itself is unchanged, including the quirks that make the output
match the old SSRS report: AllPayID = 0 means "deliberately not on payroll",
Card Number stays text so Excel does not render it as 3.53E+13, and the payroll
sheet keeps its two empty spacer columns.
"""
import csv
import io
from datetime import datetime
from typing import Dict, List, Tuple

from openpyxl import Workbook, load_workbook

REPORT_TITLE = "Monthly MealCard Adjustments for Payroll"

REQUIRED_COLUMNS = ("Card Number", "Dollars Added", "Date", "Store Name")

# Columns kept as text. Everything else is written as a number when it looks
# like one, so Excel can sum/filter it. Card Number must stay text or Excel
# renders a 14-digit card as 3.53E+13; Check No. holds values like "web".
TEXT_COLUMNS = {
    "Store Name", "Name", "Card Number", "Card Template",
    "Transaction Type", "Check No.",
}


class ReportError(ValueError):
    """The uploaded file is not a usable Paytronix detail report."""


# -- Reading the Paytronix report --------------------------------------------
def clean_card(v) -> str:
    if isinstance(v, float):
        return str(int(v))
    return str(v).strip()


def clean_amount(v) -> float:
    s = str(v).replace("$", "").replace(",", "").strip()
    try:
        return float(s) if s else 0.0
    except ValueError:
        return 0.0


def read_report(raw: bytes, filename: str = "") -> Tuple[List[str], List[list]]:
    """Read the report from bytes. Returns (header, rows), dates as datetimes."""
    name = (filename or "").lower()
    if name.endswith(".xls"):
        return _read_xls(raw)
    if name.endswith(".xlsx"):
        return _read_xlsx(raw)
    if name.endswith(".csv"):
        return _read_csv(raw)
    # No usable extension - sniff it. xls and xlsx have distinctive magic bytes.
    if raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return _read_xls(raw)
    if raw[:2] == b"PK":
        return _read_xlsx(raw)
    return _read_csv(raw)


def _read_xls(raw: bytes):
    # Imported here, not at module scope: main.py imports this module at
    # startup, and a missing xlrd would take the whole app down rather than
    # just this one file format. Exporting as .csv needs no extra package.
    try:
        import xlrd
    except ImportError:
        raise ReportError(
            "Reading the old .xls format needs the xlrd package, which is not "
            "installed on this server (pip install xlrd). In the meantime, export "
            "the report from Paytronix as .csv and upload that instead."
        )

    book = xlrd.open_workbook(file_contents=raw)
    sh = book.sheet_by_index(0)
    header = [str(sh.cell_value(0, c)).strip() for c in range(sh.ncols)]
    date_i = header.index("Date") if "Date" in header else None

    rows = []
    for r in range(1, sh.nrows):
        row = [sh.cell_value(r, c) for c in range(sh.ncols)]
        if date_i is not None and isinstance(row[date_i], float):
            row[date_i] = xlrd.xldate_as_datetime(row[date_i], book.datemode)
        rows.append(row)
    return header, rows


def _read_xlsx(raw: bytes):
    ws = load_workbook(io.BytesIO(raw), data_only=True).worksheets[0]
    header = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    rows = [
        [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1)
    ]
    return header, rows


def _read_csv(raw: bytes):
    text = raw.decode("utf-8-sig", errors="replace")
    all_rows = list(csv.reader(io.StringIO(text, newline="")))
    # The CSV export carries a meta banner line above the real header.
    hdr_i = 0
    for i, row in enumerate(all_rows[:5]):
        if any(cell.strip() == "Card Number" for cell in row):
            hdr_i = i
            break
    header = [c.strip() for c in all_rows[hdr_i]] if all_rows else []
    date_i = header.index("Date") if "Date" in header else None

    rows = []
    for row in all_rows[hdr_i + 1:]:
        if not row or not any(c.strip() for c in row):
            continue
        row = list(row)
        if date_i is not None and date_i < len(row):
            for fmt in ("%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M",
                        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
                try:
                    row[date_i] = datetime.strptime(row[date_i].strip(), fmt)
                    break
                except (ValueError, AttributeError):
                    continue
        rows.append(row)
    return header, rows


def report_date_token(filename: str) -> str:
    """Pull YYYY-MM-DD out of a ...Report-YYYYMMDD-hhmm-... filename."""
    parts = (filename or "").split("-")
    if len(parts) > 1 and len(parts[1]) == 8 and parts[1].isdigit():
        d = parts[1]
        return f"{d[:4]}-{d[4:6]}-{d[6:8]}"
    return datetime.now().strftime("%Y-%m-%d")


# -- Output 1: report with names ---------------------------------------------
def _coerce(column: str, value):
    """Match the cell types of the original hand-built report."""
    if column in TEXT_COLUMNS or isinstance(value, datetime):
        return value
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else value
    text = str(value).replace("$", "").replace(",", "").strip()
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return value
    return int(number) if number.is_integer() else number


def build_with_names(header, rows, card_names) -> bytes:
    """Raw report with the Name column filled in.

    Paytronix ships a Name column of its own, exported empty, so the usual job
    is to fill it. Older exports came without one - those get a Name column
    inserted after Store Name, where people expect to read it. Inserting
    unconditionally is what produced two Name columns side by side.
    """
    ci = {h: i for i, h in enumerate(header)}
    name_col = ci.get("Name")
    insert_at = ci.get("Store Name", 1) + 1 if name_col is None else None

    wb = Workbook()
    ws = wb.active

    out_header = list(header)
    if insert_at is not None:
        out_header.insert(insert_at, "Name")
    ws.append(out_header)

    for row in rows:
        card = clean_card(row[ci["Card Number"]])
        out = [_coerce(header[i], v) for i, v in enumerate(row)]
        if insert_at is None:
            out[name_col] = card_names.get(card, "")
        else:
            out.insert(insert_at, card_names.get(card, ""))
        ws.append(out)

    # +1 if a Name column was inserted ahead of Date, +1 again for 1-based cols
    shift = 1 if (insert_at is not None and ci["Date"] >= insert_at) else 0
    date_col = ci["Date"] + shift + 1
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, date_col)
        if isinstance(cell.value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm:ss"

    for col, width in zip("ABCDEFGHIJKLMNOP",
                          (13, 13, 30, 20, 18, 14, 18, 13, 10, 14, 16, 12, 11, 12, 11)):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -- Output 2: payroll report ------------------------------------------------
def build_payroll_rows(header, rows, card_names, by_card, by_name):
    """Group Dollars Added per employee, mirroring the SSRS proc."""
    ci = {h: i for i, h in enumerate(header)}
    totals: Dict[tuple, float] = {}
    unmapped = []

    for row in rows:
        card = clean_card(row[ci["Card Number"]])
        amount = clean_amount(row[ci["Dollars Added"]])
        name = card_names.get(card, "")

        info = by_card.get(card) or by_name.get(name)
        emp_id = (info or {}).get("employee_id", "")

        # The proc excludes AllPayID = 0 (cards deliberately not on payroll).
        if not info or not emp_id or emp_id == "0":
            unmapped.append([card, name, round(amount, 2),
                             "AllPayID = 0" if info else "no employee ID found"])
            continue

        key = (emp_id, info.get("entity", ""), info.get("name") or name)
        totals[key] = round(totals.get(key, 0.0) + amount, 2)

    def sort_key(k):
        eid = k[0]
        return (0, int(eid)) if eid.isdigit() else (1, 0)

    payroll = [[eid, amt, ent, "", "", nm]
               for (eid, ent, nm), amt in sorted(totals.items(), key=lambda kv: sort_key(kv[0]))]
    return payroll, unmapped


def build_payroll(payroll, unmapped, mapping_source: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly MealCard Adjustments fo"

    ws.append([REPORT_TITLE])
    ws.append([f"Report generated at:  {datetime.now().strftime('%m/%d/%Y %I:%M:%S %p')}"])
    ws.append([])
    ws.append(["Employee ID", "Amount", "Entity", "", "", "Name"])
    for row in payroll:
        ws.append(row)

    for col, width in zip("ABCDEF", (14, 12, 10, 3, 3, 30)):
        ws.column_dimensions[col].width = width
    for r in range(5, ws.max_row + 1):
        ws.cell(r, 2).number_format = "0.00"

    if unmapped:
        ws2 = wb.create_sheet("Unmapped")
        ws2.append(["Card Number", "Name", "Dollars Added", "Reason"])
        ws2.append([f"NOT included in the payroll sheet. Mapping source: {mapping_source}"])
        ws2.append([])
        for row in unmapped:
            ws2.append(row)
        for col, width in zip("ABCD", (20, 32, 15, 60)):
            ws2.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -- The whole job -----------------------------------------------------------
def build_reports(raw: bytes, filename: str, card_names: Dict[str, str],
                  by_card: dict, by_name: dict, mapping_source: str) -> dict:
    """Read an uploaded report and return both workbooks plus what to show.

    Raises ReportError if the upload is not the report we expect - the message
    goes straight to the user, so it says which column is missing rather than
    "KeyError".
    """
    try:
        header, rows = read_report(raw, filename)
    except ReportError:
        raise
    except Exception as e:
        raise ReportError(
            f"Could not read that file as a Paytronix report ({e.__class__.__name__}). "
            "Export it from Paytronix as .xls or .csv and upload that."
        ) from e

    ci = {h: i for i, h in enumerate(header)}
    missing = [c for c in REQUIRED_COLUMNS if c not in ci]
    if missing:
        raise ReportError(
            f"That file is missing the {', '.join(repr(m) for m in missing)} column"
            f"{'s' if len(missing) > 1 else ''}. "
            f"Columns found: {', '.join(header) if header else '(none)'}"
        )
    if not rows:
        raise ReportError("That report has a header but no data rows.")

    token = report_date_token(filename)

    names_xlsx = build_with_names(header, rows, card_names)
    payroll_rows, unmapped = build_payroll_rows(header, rows, card_names, by_card, by_name)
    payroll_xlsx = build_payroll(payroll_rows, unmapped, mapping_source)

    dates = [r[ci["Date"]] for r in rows if isinstance(r[ci["Date"]], datetime)]
    missing_names = sorted({
        clean_card(r[ci["Card Number"]]) for r in rows
        if clean_card(r[ci["Card Number"]]) not in card_names
    })

    return {
        "token": token,
        "names_filename": f"StoredValueReport_with_names_{token}.xlsx",
        "names_xlsx": names_xlsx,
        "payroll_filename": f"{REPORT_TITLE}_{token}.xlsx",
        "payroll_xlsx": payroll_xlsx,
        "stats": {
            "rows": len(rows),
            "date_from": min(dates).strftime("%Y-%m-%d %H:%M") if dates else None,
            "date_to": max(dates).strftime("%Y-%m-%d %H:%M") if dates else None,
            "total_added": round(sum(clean_amount(r[ci["Dollars Added"]]) for r in rows), 2),
            "employees": len(payroll_rows),
            "payroll_total": round(sum(r[1] for r in payroll_rows), 2),
            "mapping_source": mapping_source,
            "cards_without_names": missing_names,
        },
        "unmapped": [
            {"card_number": c, "name": n, "amount": a, "reason": why}
            for c, n, a, why in unmapped
        ],
    }
