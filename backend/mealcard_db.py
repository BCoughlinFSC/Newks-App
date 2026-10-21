"""
mealcard_db.py - SQLite store for the monthly meal card reload.

Replaces the two files the desktop scripts used to read:

  * "Gift Card _ Meal Card Numbers - Support Center 2026.xlsx" - the card
    roster plus a tick per period. Now mealcard_cards + mealcard_periods.
  * MealCardsToAllPay.csv, exported by hand from dev_aloha - the card to
    employee-ID mapping the payroll report joins to. Now mealcard_employee_map.

The reload bot runs on someone's PC (it drives their logged-in Paytronix
Chrome), so it asks this app for its worklist and posts each card's outcome
back. That makes these tables the record of what was loaded, instead of a
spreadsheet that has to be open on one person's machine.

See main.py /api/mealcard/* and backend/mealcard_bot/.
"""
import csv
import io
import json
import logging
import os
import re
import sqlite3
from datetime import datetime, timezone
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MEALCARD_DB_PATH = os.getenv(
    "MEALCARD_DB_PATH",
    os.path.join(os.path.dirname(__file__), "data", "mealcard", "mealcard.db"),
)

DEFAULT_TARGET = 160.00

# The card list the app ships with, so a new deployment works straight away
# instead of starting empty and waiting for someone to upload a spreadsheet.
SEED_PATH = os.getenv(
    "MEALCARD_SEED_PATH",
    os.path.join(os.path.dirname(__file__), "data", "mealcard_seed.json"),
)

logger = logging.getLogger(__name__)
_seed_checked = False

_SCHEMA = """
CREATE TABLE IF NOT EXISTS mealcard_cards (
    card_number   TEXT PRIMARY KEY,
    name          TEXT NOT NULL,
    target_amount REAL,                      -- NULL means "use the default"
    active        INTEGER NOT NULL DEFAULT 1,
    sort_order    INTEGER NOT NULL DEFAULT 0,
    notes         TEXT,
    updated_at    TEXT
);

CREATE TABLE IF NOT EXISTS mealcard_periods (
    card_number    TEXT    NOT NULL,
    year           INTEGER NOT NULL,
    period         INTEGER NOT NULL,
    status         TEXT    NOT NULL,         -- done | skipped | error
    amount         REAL,                     -- what the bot added
    balance_before REAL,
    detail         TEXT,
    run_id         INTEGER,
    recorded_at    TEXT    NOT NULL,
    PRIMARY KEY (card_number, year, period)
);

CREATE TABLE IF NOT EXISTS mealcard_employee_map (
    card_number TEXT PRIMARY KEY,
    employee_id TEXT,
    entity      TEXT,
    name        TEXT,
    source      TEXT,
    updated_at  TEXT
);

-- Weaker name-matched fallback, used only for cards the card map misses.
CREATE TABLE IF NOT EXISTS mealcard_employee_map_by_name (
    name        TEXT PRIMARY KEY,
    employee_id TEXT,
    entity      TEXT,
    source      TEXT,
    updated_at  TEXT
);

CREATE TABLE IF NOT EXISTS mealcard_runs (
    id           INTEGER PRIMARY KEY AUTOINCREMENT,
    year         INTEGER NOT NULL,
    period       INTEGER NOT NULL,
    -- Set when the run STARTS, not when it ends: while a run is in flight,
    -- "is this spending money?" is the question people actually have.
    mode         TEXT    NOT NULL DEFAULT 'live',     -- live | test
    status       TEXT    NOT NULL DEFAULT 'running',  -- running | finished | abandoned
    started_at   TEXT    NOT NULL,
    finished_at  TEXT,
    requested_by TEXT,
    completed    INTEGER NOT NULL DEFAULT 0,
    skipped      INTEGER NOT NULL DEFAULT 0,
    errors       INTEGER NOT NULL DEFAULT 0,
    log          TEXT
);

CREATE INDEX IF NOT EXISTS idx_mealcard_periods_year ON mealcard_periods(year, period);
CREATE INDEX IF NOT EXISTS idx_mealcard_runs_started ON mealcard_runs(started_at);
"""


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def get_connection(db_path: str = MEALCARD_DB_PATH) -> sqlite3.Connection:
    os.makedirs(os.path.dirname(db_path), exist_ok=True)
    conn = sqlite3.connect(db_path, check_same_thread=False, timeout=30)
    conn.row_factory = sqlite3.Row
    return conn


def initialize_db(db_path: str = MEALCARD_DB_PATH) -> None:
    global _seed_checked
    with get_connection(db_path) as conn:
        conn.executescript(_SCHEMA)
        try:
            conn.execute("ALTER TABLE mealcard_runs ADD COLUMN mode TEXT NOT NULL DEFAULT 'live'")
        except sqlite3.OperationalError:
            pass    # already there
        # Runs from before `mode` existed recorded a test by setting the status
        # to dry-run at the end. Move those across so one idea lives in one column.
        conn.execute("UPDATE mealcard_runs SET mode = 'test', status = 'finished' "
                     "WHERE status = 'dry-run'")
        conn.commit()

    # First touch in this process: fill empty tables from the data the app
    # ships with, so a fresh deployment is usable without importing anything.
    if not _seed_checked:
        _seed_checked = True
        try:
            result = load_seed(db_path)
            if result.get("seeded"):
                logger.info("mealcard: seeded %s cards, %s ticks, %s employee names "
                            "from the shipped data",
                            result.get("cards"), result.get("ticks"), result.get("names"))
        except Exception as exc:          # never block the app over seed data
            logger.warning("mealcard: could not load the shipped card data (%s: %s)",
                           exc.__class__.__name__, exc)


def clean_card(v) -> str:
    """Card numbers arrive as floats from Excel and as strings everywhere else."""
    if isinstance(v, float):
        return str(int(v))
    return str(v).strip()


def parse_target_amount(name: str) -> float:
    """Pull a custom amount out of a name like 'Cynthia Baxter (Board Meeting $600.00)'.

    The roster carried per-card exceptions in the name text, so seeding reads
    them the same way the desktop bot did. Once a card is in the table its
    target_amount is editable on its own.
    """
    match = re.search(r"\$(\d+(?:\.\d{2})?)", name or "")
    return float(match.group(1)) if match else DEFAULT_TARGET


# -- Cards --------------------------------------------------------------------
def list_cards(active_only: bool = True, db_path: str = MEALCARD_DB_PATH) -> List[Dict]:
    initialize_db(db_path)
    sql = "SELECT * FROM mealcard_cards"
    if active_only:
        sql += " WHERE active = 1"
    sql += " ORDER BY sort_order, name"
    with get_connection(db_path) as conn:
        return [dict(r) for r in conn.execute(sql)]


def upsert_card(card_number: str, name: str, target_amount: Optional[float] = None,
                active: bool = True, sort_order: int = 0, notes: str = "",
                db_path: str = MEALCARD_DB_PATH) -> None:
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        conn.execute(
            """INSERT INTO mealcard_cards
                   (card_number, name, target_amount, active, sort_order, notes, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?)
               ON CONFLICT(card_number) DO UPDATE SET
                   name = excluded.name,
                   target_amount = excluded.target_amount,
                   active = excluded.active,
                   sort_order = excluded.sort_order,
                   notes = excluded.notes,
                   updated_at = excluded.updated_at""",
            (clean_card(card_number), name.strip(), target_amount,
             1 if active else 0, sort_order, notes, _now()),
        )
        conn.commit()


def card_names(db_path: str = MEALCARD_DB_PATH) -> Dict[str, str]:
    """Card number -> name, for every card including inactive ones.

    The payroll report looks up names for whatever Paytronix reported, which
    can include a card retired mid-month, so this deliberately ignores `active`.
    """
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        return {r["card_number"]: r["name"]
                for r in conn.execute("SELECT card_number, name FROM mealcard_cards")}


def target_for(card: Dict) -> float:
    t = card.get("target_amount")
    return float(t) if t is not None else DEFAULT_TARGET


# -- Importing the roster workbook -------------------------------------------
# Column A is the name, column B the card number, and each period gets its own
# column from C on - so period N lives in column 2 + N.
FIRST_PERIOD_COL = 3


class RosterError(ValueError):
    """The workbook is not the gift card roster we expect."""


# A period column reads "Period 01 Meal Card " on the original spreadsheet and
# "Period 01" on the one this app exports. Both say the number the same way.
PERIOD_HEADER_RE = re.compile(r"^period\s*0*(\d{1,2})\b", re.I)

# What a tick can look like. The original sheet stores booleans; this app's
# export writes the words, and a round trip has to survive either.
TICK_WORDS = {
    "loaded": "done", "done": "done", "true": "done", "yes": "done", "x": "done",
    "skipped (at target)": "skipped", "skipped": "skipped",
    "error": "error",
}


def _pick_sheet(wb, wanted: str):
    """Find the roster tab, forgiving the name where there is no ambiguity.

    The original workbook names its tabs by year ("2026"); the export this app
    produces names its one "Meal Cards 2026". Someone editing the download and
    sending it back should not have to rename a tab to get it accepted.
    """
    if wanted in wb.sheetnames:
        return wb[wanted]
    loose = [n for n in wb.sheetnames if wanted and wanted.lower() in n.lower()]
    if len(loose) == 1:
        return wb[loose[0]]
    if len(wb.sheetnames) == 1:
        return wb[wb.sheetnames[0]]
    raise RosterError(
        f"That workbook has no {wanted!r} sheet. It has: {', '.join(wb.sheetnames)}")


def _header_columns(ws) -> Dict[str, int]:
    out = {}
    for c in range(1, ws.max_column + 1):
        label = str(ws.cell(1, c).value or "").strip()
        if label:
            out.setdefault(label.lower(), c)
    return out


def parse_roster(source, sheet: str) -> List[Dict]:
    """Read a roster workbook. `source` is a path or the file's bytes.

    Reads by column heading rather than by position, so it takes both the
    original gift card spreadsheet (Name, Card Number, then a column per
    period) and the sheet this app exports, which adds Target Amount and
    Active in between. That means the download can be edited and sent back,
    which is how most changes will actually be made.

    Parsing is kept separate from writing so a caller can show what an import
    would do before it does it.
    """
    wb = load_workbook(
        io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else source,
        data_only=True,
    )
    try:
        ws = _pick_sheet(wb, sheet)
        headers = _header_columns(ws)

        name_col = headers.get("name")
        card_col = next((c for label, c in headers.items()
                         if "card number" in label or label == "number"
                         or label == "card"), None)
        if not name_col or not card_col:
            found = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
            raise RosterError(
                "That sheet does not look like a card list - it needs a Name column and a "
                f"Card Number column. Its headings are: {', '.join(h for h in found if h)}")

        target_col = headers.get("target amount") or headers.get("target")
        active_col = headers.get("active")
        period_cols = {}
        for label, c in headers.items():
            m = PERIOD_HEADER_RE.match(label)
            if m and 1 <= int(m.group(1)) <= 13:
                period_cols.setdefault(int(m.group(1)), c)

        rows = []
        for r in range(2, ws.max_row + 1):
            name_cell, card_cell = ws.cell(r, name_col).value, ws.cell(r, card_col).value
            if not name_cell or not card_cell:
                continue
            name = str(name_cell).strip()

            # An explicit Target Amount column wins; otherwise fall back to the
            # amount written into the name, which is how the original sheet
            # carried its exceptions.
            target = parse_target_amount(name)
            if target_col:
                raw = ws.cell(r, target_col).value
                if isinstance(raw, (int, float)):
                    target = float(raw)
                elif raw not in (None, ""):
                    try:
                        target = float(str(raw).replace("$", "").replace(",", "").strip())
                    except ValueError:
                        pass

            active = True
            if active_col:
                active = str(ws.cell(r, active_col).value or "yes").strip().lower() \
                    not in ("no", "false", "0", "inactive")

            periods = {}
            for period, col in period_cols.items():
                value = ws.cell(r, col).value
                if value is True:
                    periods[period] = "done"
                elif isinstance(value, str) and value.strip():
                    status = TICK_WORDS.get(value.strip().lower())
                    if status:
                        periods[period] = status

            rows.append({
                "card_number": clean_card(card_cell),
                "name": name,
                "target_amount": target,
                "active": active,
                "sort_order": r,
                "periods": periods,
            })
    finally:
        wb.close()

    if not rows:
        raise RosterError(f"No cards found on sheet {ws.title!r}.")
    return rows


def load_seed(db_path: str = MEALCARD_DB_PATH, path: str = SEED_PATH) -> Dict:
    """Fill empty tables from the card data the app ships with.

    Only empty ones. A database someone has since edited must never be
    overwritten by whatever the repo happens to ship, so this fills a blank
    install and then stays out of the way for good. Deleting the seed file
    changes nothing for an app that has already run.
    """
    if not os.path.exists(path):
        return {"seeded": False, "reason": "no seed file shipped"}

    with open(path, encoding="utf-8") as fh:
        seed = json.load(fh)

    with get_connection(db_path) as conn:
        have_cards = conn.execute("SELECT COUNT(*) FROM mealcard_cards").fetchone()[0]
        have_names = conn.execute(
            "SELECT COUNT(*) FROM mealcard_employee_map_by_name").fetchone()[0]
        have_map = conn.execute("SELECT COUNT(*) FROM mealcard_employee_map").fetchone()[0]

    cards = ticks = names = 0
    year = int(seed.get("year") or datetime.now().year)

    if not have_cards:
        rows = [{
            "card_number": clean_card(c["card_number"]),
            "name": c["name"],
            # The file stores null for "use the default"; import_roster takes a
            # real number and puts the null back.
            "target_amount": (c["target_amount"] if c.get("target_amount") is not None
                              else DEFAULT_TARGET),
            "sort_order": c.get("sort_order", 0),
            "periods": c.get("periods", []),
        } for c in seed.get("cards", [])]
        if rows:
            # No deactivation on a seed: the tables were empty, so there is
            # nothing it could possibly be retiring.
            written = import_roster(rows, year, source="the list shipped with the app",
                                    db_path=db_path)
            cards, ticks = written["cards"], written["ticks"]

    # Only when there is no mapping of either kind: once a real export has been
    # imported, the shipped name-matched fallback has nothing to add.
    if not have_names and not have_map and seed.get("employee_names"):
        names = upsert_employee_names(
            seed["employee_names"],
            source=seed.get("employee_names_source", "shipped with the app"),
            db_path=db_path)

    return {"seeded": bool(cards or names), "cards": cards, "ticks": ticks,
            "names": names, "year": year}


STATUS_WORDS = {"done": "Loaded", "skipped": "Skipped (at target)", "error": "Error"}


def roster_xlsx(year: int, include_inactive: bool = False,
                db_path: str = MEALCARD_DB_PATH) -> bytes:
    """The card list as a workbook: who, which card, what it tops up to.

    Card Number is written as text. It is a 14-digit number, and Excel renders
    a long number in General format as 3.53E+13 - which is what the roster
    spreadsheet this replaces does.
    """
    cards = list_cards(active_only=not include_inactive, db_path=db_path)
    grid = year_grid(year, db_path=db_path)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Meal Cards {year}"

    header = ["Name", "Card Number", "Target Amount", "Active"]
    header += [f"Period {p:02d}" for p in range(1, 13)]
    ws.append(header)

    for card in cards:
        periods = grid.get(card["card_number"], {})
        ws.append([
            card["name"],
            card["card_number"],
            target_for(card),
            "Yes" if card["active"] else "No",
            *[STATUS_WORDS.get(periods.get(p), "") for p in range(1, 13)],
        ])

    head_fill = PatternFill("solid", fgColor="C41E3A")
    head_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(header) + 1):
        cell = ws.cell(1, col)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrapText=True)

    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).number_format = "@"              # keep every digit
        ws.cell(row, 3).number_format = '"$"#,##0.00'
        ws.cell(row, 4).alignment = Alignment(horizontal="center")
        for col in range(5, len(header) + 1):
            ws.cell(row, col).alignment = Alignment(horizontal="center")

    for col, width in enumerate([34, 20, 15, 9] + [13] * 12, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"                               # name and card stay put
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{ws.max_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def cards_missing_from(rows: List[Dict], db_path: str = MEALCARD_DB_PATH) -> List[Dict]:
    """Active cards the app holds that this sheet does not mention.

    Almost always someone who has left: they were deleted from the spreadsheet,
    and the import is how that is meant to reach the app.
    """
    incoming = {clean_card(r["card_number"]) for r in rows}
    return [c for c in list_cards(active_only=True, db_path=db_path)
            if c["card_number"] not in incoming]


def import_roster(rows: List[Dict], year: int, source: str = "",
                  deactivate_missing: bool = False,
                  db_path: str = MEALCARD_DB_PATH) -> Dict:
    """Write parsed roster rows in. Cards upsert, ticks upsert per period.

    With deactivate_missing, a card the sheet no longer lists is marked
    inactive: it drops out of the roster and the bot's worklist, but keeps its
    row and its period history, so last year's record still reads correctly and
    the decision is reversible. Nothing is ever deleted, and the caller is
    expected to have shown the names first - see the dry run in
    /api/mealcard/roster-import.
    """
    initialize_db(db_path)
    missing = cards_missing_from(rows, db_path) if deactivate_missing else []
    cards = ticks = 0
    for row in rows:
        upsert_card(
            card_number=row["card_number"],
            name=row["name"],
            # Only store a target when it differs from the default; the rest
            # stay NULL so changing the default moves them all.
            target_amount=(row["target_amount"]
                           if row["target_amount"] != DEFAULT_TARGET else None),
            active=row.get("active", True),
            sort_order=row["sort_order"],
            db_path=db_path,
        )
        cards += 1
        # A list of period numbers (the shipped seed) or {period: status} (a
        # parsed workbook, which can say skipped and errored too).
        periods = row.get("periods") or {}
        if not isinstance(periods, dict):
            periods = {p: "done" for p in periods}
        for period, status in periods.items():
            record_period(row["card_number"], year, int(period), status or "done",
                          detail=f"imported from {source}" if source else "imported",
                          db_path=db_path)
            ticks += 1

    for card in missing:
        upsert_card(card_number=card["card_number"], name=card["name"],
                    target_amount=card["target_amount"], active=False,
                    sort_order=card["sort_order"], notes=card.get("notes") or "",
                    db_path=db_path)

    return {"cards": cards, "ticks": ticks, "year": year,
            "deactivated": [c["name"] for c in missing]}


# -- Period completion --------------------------------------------------------
def period_status(year: int, period: int, db_path: str = MEALCARD_DB_PATH) -> Dict[str, Dict]:
    """card_number -> that card's row for this period, if it has one."""
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        rows = conn.execute(
            "SELECT * FROM mealcard_periods WHERE year = ? AND period = ?",
            (year, period),
        )
        return {r["card_number"]: dict(r) for r in rows}


def year_grid(year: int, db_path: str = MEALCARD_DB_PATH) -> Dict[str, Dict[int, str]]:
    """card_number -> {period: status} for a whole year, for the roster table."""
    initialize_db(db_path)
    grid: Dict[str, Dict[int, str]] = {}
    with get_connection(db_path) as conn:
        for r in conn.execute(
            "SELECT card_number, period, status FROM mealcard_periods WHERE year = ?",
            (year,),
        ):
            grid.setdefault(r["card_number"], {})[r["period"]] = r["status"]
    return grid


def record_period(card_number: str, year: int, period: int, status: str,
                  amount: Optional[float] = None, balance_before: Optional[float] = None,
                  detail: str = "", run_id: Optional[int] = None,
                  db_path: str = MEALCARD_DB_PATH) -> None:
    """Write one card's outcome for a period. Re-running a card overwrites it."""
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        conn.execute(
            """INSERT INTO mealcard_periods
                   (card_number, year, period, status, amount, balance_before,
                    detail, run_id, recorded_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
               ON CONFLICT(card_number, year, period) DO UPDATE SET
                   status = excluded.status,
                   amount = excluded.amount,
                   balance_before = excluded.balance_before,
                   detail = excluded.detail,
                   run_id = excluded.run_id,
                   recorded_at = excluded.recorded_at""",
            (clean_card(card_number), year, period, status, amount, balance_before,
             detail, run_id, _now()),
        )
        conn.commit()


def clear_period(card_number: str, year: int, period: int,
                 db_path: str = MEALCARD_DB_PATH) -> bool:
    """Un-tick one card for a period, so the next run picks it up again."""
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        cur = conn.execute(
            "DELETE FROM mealcard_periods WHERE card_number = ? AND year = ? AND period = ?",
            (clean_card(card_number), year, period),
        )
        conn.commit()
        return cur.rowcount > 0


# -- Employee mapping ---------------------------------------------------------
def replace_employee_map(rows: List[Dict], source: str,
                         db_path: str = MEALCARD_DB_PATH) -> int:
    """Replace the card -> employee mapping wholesale.

    A fresh export from dev_aloha.dbo.MealCardsToAllPay is the whole truth, so
    importing one drops what was there. Merging would keep resurrecting cards
    that the export deliberately no longer has.
    """
    initialize_db(db_path)
    stamped = _now()
    with get_connection(db_path) as conn:
        conn.execute("DELETE FROM mealcard_employee_map")
        conn.executemany(
            """INSERT INTO mealcard_employee_map
                   (card_number, employee_id, entity, name, source, updated_at)
               VALUES (?, ?, ?, ?, ?, ?)""",
            [(clean_card(r.get("card_number", "")), str(r.get("employee_id", "") or ""),
              str(r.get("entity", "") or ""), str(r.get("name", "") or ""), source, stamped)
             for r in rows if clean_card(r.get("card_number", ""))],
        )
        conn.commit()
        return conn.execute("SELECT COUNT(*) FROM mealcard_employee_map").fetchone()[0]


def upsert_employee_names(rows: List[Dict], source: str,
                          db_path: str = MEALCARD_DB_PATH) -> int:
    """Merge into the name-matched fallback map.

    Unlike the card map this merges rather than replaces: each seed covers only
    the people reloaded that month, so replacing would shrink it every time.
    """
    initialize_db(db_path)
    stamped = _now()
    with get_connection(db_path) as conn:
        conn.executemany(
            """INSERT INTO mealcard_employee_map_by_name
                   (name, employee_id, entity, source, updated_at)
               VALUES (?, ?, ?, ?, ?)
               ON CONFLICT(name) DO UPDATE SET
                   employee_id = excluded.employee_id,
                   entity = excluded.entity,
                   source = excluded.source,
                   updated_at = excluded.updated_at""",
            [(str(r.get("name", "")).strip(), str(r.get("employee_id", "") or ""),
              str(r.get("entity", "") or ""), source, stamped)
             for r in rows if str(r.get("name", "")).strip()],
        )
        conn.commit()
        return conn.execute("SELECT COUNT(*) FROM mealcard_employee_map_by_name").fetchone()[0]


def employee_map(db_path: str = MEALCARD_DB_PATH) -> Tuple[Dict, Dict, str]:
    """(by_card, by_name, description) - the same shape the desktop script used.

    by_card is the authoritative dev_aloha mapping. by_name is the weaker
    fallback lifted from past payroll exports, used only where by_card misses.
    """
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        by_card = {
            r["card_number"]: {"employee_id": r["employee_id"] or "",
                               "entity": r["entity"] or "",
                               "name": r["name"] or ""}
            for r in conn.execute("SELECT * FROM mealcard_employee_map")
        }
        by_name = {
            r["name"]: {"employee_id": r["employee_id"] or "",
                        "entity": r["entity"] or "",
                        "name": r["name"]}
            for r in conn.execute("SELECT * FROM mealcard_employee_map_by_name")
        }
        card_src = conn.execute(
            "SELECT source FROM mealcard_employee_map LIMIT 1").fetchone()
        name_src = conn.execute(
            "SELECT source FROM mealcard_employee_map_by_name LIMIT 1").fetchone()

    if by_card:
        # Describe it rather than grading it: once IDs can be set one at a time,
        # "authoritative" would be claimed by a single hand-typed row.
        desc = f"{len(by_card)} from {card_src['source'] if card_src else 'the card mapping'}"
        if by_name:
            desc += f", {len(by_name)} name-matched from earlier payroll exports"
        return by_card, by_name, desc
    if by_name:
        return {}, by_name, (f"{len(by_name)} names from "
                             f"{name_src['source'] if name_src else 'payroll exports'} "
                             "- FALLBACK, name-matched")
    return {}, {}, "NONE - no mapping imported yet"


MAPPING_COLUMNS = ["CardNumber", "EmployeeID", "Entity", "Name"]


def set_employee_for_card(card_number: str, employee_id: str, entity: str = "",
                          source: str = "set on the Meal Cards page",
                          db_path: str = MEALCARD_DB_PATH) -> Dict:
    """Give one card an employee ID, without disturbing the other 33.

    An empty id clears it, from the name-matched fallback as well as the card
    map - otherwise clearing a card would just fall back to the older guess and
    look like nothing happened. "0" is kept, not cleared: that is the deliberate
    "never on the payroll report" marker, not a blank.
    """
    initialize_db(db_path)
    card_number = clean_card(card_number)
    stamped = _now()

    with get_connection(db_path) as conn:
        row = conn.execute("SELECT name FROM mealcard_cards WHERE card_number = ?",
                           (card_number,)).fetchone()
        if not row:
            return {"found": False}
        name = row["name"]

        employee_id = str(employee_id or "").strip()
        if employee_id == "":
            conn.execute("DELETE FROM mealcard_employee_map WHERE card_number = ?",
                         (card_number,))
            conn.execute("DELETE FROM mealcard_employee_map_by_name WHERE name = ?", (name,))
        else:
            conn.execute(
                """INSERT INTO mealcard_employee_map
                       (card_number, employee_id, entity, name, source, updated_at)
                   VALUES (?, ?, ?, ?, ?, ?)
                   ON CONFLICT(card_number) DO UPDATE SET
                       employee_id = excluded.employee_id,
                       entity = excluded.entity,
                       name = excluded.name,
                       source = excluded.source,
                       updated_at = excluded.updated_at""",
                (card_number, employee_id, entity.strip(), name, source, stamped),
            )
        conn.commit()

    return {"found": True, "card_number": card_number, "name": name,
            "employee_id": employee_id, "entity": entity.strip()}


def employee_map_csv(db_path: str = MEALCARD_DB_PATH) -> bytes:
    """The current mapping, in the same shape MealCardsToAllPay.csv arrives in.

    One row per active card, including the ones with no employee ID yet, so the
    download doubles as the template for fixing it: open it, fill the blanks in
    the EmployeeID column, upload it back. That means nobody has to know the
    format, and nobody has to reproduce the rows that are already right.
    """
    by_card, by_name, _ = employee_map(db_path)

    buf = io.StringIO(newline="")
    writer = csv.DictWriter(buf, fieldnames=MAPPING_COLUMNS, lineterminator="\r\n")
    writer.writeheader()
    for card in list_cards(active_only=True, db_path=db_path):
        info = by_card.get(card["card_number"]) or by_name.get(card["name"]) or {}
        writer.writerow({
            "CardNumber": card["card_number"],
            "EmployeeID": info.get("employee_id", ""),
            "Entity": info.get("entity", ""),
            # Prefer the payroll spelling where we have one; fall back to the
            # roster so a blank row still says who it is for.
            "Name": info.get("name") or card["name"],
        })
    # utf-8-sig: Excel reads a bare utf-8 CSV as the system codepage and
    # mangles any name with an accent or apostrophe in it.
    return buf.getvalue().encode("utf-8-sig")


def mapping_summary(db_path: str = MEALCARD_DB_PATH) -> Dict:
    """How the mapping stands, including how much of the roster it actually covers.

    The count is the useful number: a mapping can exist and still miss most
    people, and what that costs is names landing on the Unmapped tab at payroll
    time rather than on the sheet.
    """
    by_card, by_name, desc = employee_map(db_path)
    cards = list_cards(active_only=True, db_path=db_path)
    covered = off_payroll = 0
    unknown = []
    for card in cards:
        info = by_card.get(card["card_number"]) or by_name.get(card["name"]) or {}
        emp_id = info.get("employee_id") or ""
        if emp_id == "0":
            # Deliberately off payroll - a corporate expense, not a missing id.
            off_payroll += 1
        elif emp_id:
            covered += 1
        else:
            unknown.append(card["name"])
    return {"by_card": len(by_card), "by_name": len(by_name),
            "description": desc, "authoritative": bool(by_card),
            "covers": covered, "of": len(cards),
            "off_payroll": off_payroll,
            "unknown": unknown, "unknown_count": len(unknown)}


# -- Runs ---------------------------------------------------------------------
def start_run(year: int, period: int, requested_by: str = "", dry_run: bool = False,
              db_path: str = MEALCARD_DB_PATH) -> int:
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        cur = conn.execute(
            """INSERT INTO mealcard_runs (year, period, mode, status, started_at, requested_by)
               VALUES (?, ?, ?, 'running', ?, ?)""",
            (year, period, "test" if dry_run else "live", _now(), requested_by),
        )
        conn.commit()
        return cur.lastrowid


def finish_run(run_id: int, completed: int = 0, skipped: int = 0, errors: int = 0,
               log: str = "", status: str = "finished",
               db_path: str = MEALCARD_DB_PATH) -> None:
    initialize_db(db_path)
    # An older copy of the bot reports a test by finishing with status
    # "dry-run" and says nothing at start. Fold that back into the mode column
    # so one idea has one spelling, and so a test from an old bot is still
    # labelled a test rather than passing for a live run.
    was_dry = status == "dry-run"
    if was_dry:
        status = "finished"
    with get_connection(db_path) as conn:
        if was_dry:
            conn.execute("UPDATE mealcard_runs SET mode = 'test' WHERE id = ?", (run_id,))
        conn.execute(
            """UPDATE mealcard_runs
                  SET status = ?, finished_at = ?, completed = ?, skipped = ?,
                      errors = ?, log = ?
                WHERE id = ?""",
            (status, _now(), completed, skipped, errors, log, run_id),
        )
        conn.commit()


def abandon_run(run_id: int, db_path: str = MEALCARD_DB_PATH) -> bool:
    """Close out a run that stopped reporting, without inventing an outcome.

    Keeps whatever counts it managed to record. Only touches a run still marked
    running, so this can never overwrite a finished one.
    """
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        cur = conn.execute(
            "UPDATE mealcard_runs SET status = 'abandoned', finished_at = ? "
            "WHERE id = ? AND status = 'running'",
            (_now(), run_id),
        )
        conn.commit()
        return cur.rowcount > 0


def append_run_log(run_id: int, text: str, db_path: str = MEALCARD_DB_PATH) -> None:
    """Append a line to a running job's log, so the page can show progress."""
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        conn.execute(
            "UPDATE mealcard_runs SET log = COALESCE(log, '') || ? WHERE id = ?",
            (text if text.endswith("\n") else text + "\n", run_id),
        )
        conn.commit()


def get_run(run_id: int, db_path: str = MEALCARD_DB_PATH) -> Optional[Dict]:
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        row = conn.execute("SELECT * FROM mealcard_runs WHERE id = ?", (run_id,)).fetchone()
        return dict(row) if row else None


def list_runs(limit: int = 25, db_path: str = MEALCARD_DB_PATH) -> List[Dict]:
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        rows = conn.execute(
            """SELECT id, year, period, mode, status, started_at, finished_at, requested_by,
                      completed, skipped, errors
                 FROM mealcard_runs ORDER BY id DESC LIMIT ?""",
            (limit,),
        )
        return [dict(r) for r in rows]


def set_card_active(card_number: str, active: bool,
                    db_path: str = MEALCARD_DB_PATH) -> bool:
    """Retire a card, or bring one back. Never deletes: the period history is
    the record of what was loaded and has to survive someone leaving."""
    initialize_db(db_path)
    with get_connection(db_path) as conn:
        cur = conn.execute(
            "UPDATE mealcard_cards SET active = ?, updated_at = ? WHERE card_number = ?",
            (1 if active else 0, _now(), clean_card(card_number)),
        )
        conn.commit()
        return cur.rowcount > 0


# -- The worklist the bot asks for -------------------------------------------
def worklist(year: int, period: int, include_done: bool = False,
             db_path: str = MEALCARD_DB_PATH) -> List[Dict]:
    """Active cards for this period, with the target and what already happened.

    Cards already marked done are left out unless include_done, which is what
    stops a second run of the month from loading everyone twice.
    """
    done = period_status(year, period, db_path)
    out = []
    for card in list_cards(active_only=True, db_path=db_path):
        prior = done.get(card["card_number"])
        if prior and prior["status"] in ("done", "skipped") and not include_done:
            continue
        out.append({
            "card_number": card["card_number"],
            "name": card["name"],
            "target_amount": target_for(card),
            "already": prior["status"] if prior else None,
        })
    return out
